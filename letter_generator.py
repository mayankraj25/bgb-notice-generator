"""Build Extension and Encashment Notice .docx files using python-docx.

Formatting matches the reference PDF exactly:
- Times New Roman 12pt throughout
- A4 page, 1-inch margins
- Justified body paragraphs
- Bold/underline applied per PDF reference
- Indian comma format for encashment amounts (8,25,000.00)
"""

import os
from datetime import date

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from amount_to_words import amount_to_words
from utils import (
    format_date, extended_validity, submission_deadline,
    build_output_filename, indian_amount_format,
)

TNR            = "Times New Roman"
FONT_SIZE      = Pt(12)
SPACE_AFTER    = Pt(6)
SPACE_BEFORE   = Pt(0)
JUSTIFY        = WD_ALIGN_PARAGRAPH.JUSTIFY
LEFT           = WD_ALIGN_PARAGRAPH.LEFT
CENTER         = WD_ALIGN_PARAGRAPH.CENTER


# ---------------------------------------------------------------------------
# Document factory
# ---------------------------------------------------------------------------

def _new_doc() -> Document:
    doc = Document()
    for section in doc.sections:
        section.page_height   = Cm(29.7)
        section.page_width    = Cm(21.0)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
    style = doc.styles['Normal']
    style.font.name            = TNR
    style.font.size            = FONT_SIZE
    style.paragraph_format.space_after  = SPACE_AFTER
    style.paragraph_format.space_before = SPACE_BEFORE
    style.paragraph_format.line_spacing = Pt(12)
    return doc


# ---------------------------------------------------------------------------
# Run helpers
# ---------------------------------------------------------------------------

def _run(para, text: str, bold=False, underline=False):
    """Add a run with optional bold/underline to an existing paragraph."""
    r = para.add_run(text)
    r.font.name    = TNR
    r.font.size    = FONT_SIZE
    r.bold         = bold
    r.underline    = underline
    return r


def _para(doc, text="", bold=False, underline=False,
          align=LEFT, space_after=SPACE_AFTER):
    """Add a simple single-run paragraph."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after  = space_after
    p.paragraph_format.space_before = SPACE_BEFORE
    if text:
        _run(p, text, bold=bold, underline=underline)
    return p


def _mixed(doc, segments, align=JUSTIFY, space_after=SPACE_AFTER,
           left_indent=None):
    """
    Add a paragraph with mixed formatting.
    segments: list of (text, bold, underline) tuples.
    """
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after  = space_after
    p.paragraph_format.space_before = SPACE_BEFORE
    if left_indent is not None:
        p.paragraph_format.left_indent = left_indent
    for item in segments:
        if len(item) == 2:
            text, bold = item
            underline = False
        else:
            text, bold, underline = item
        _run(p, text, bold=bold, underline=underline)
    return p


# ---------------------------------------------------------------------------
# Table helpers
# ---------------------------------------------------------------------------

def _remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'none')
        tblBorders.append(el)
    tblPr.append(tblBorders)


def _no_border_table(doc, rows, cols):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    _remove_table_borders(t)
    return t


def _bordered_table(doc, rows, cols):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = 'Table Grid'
    return t


def _cell(cell, segments, align=LEFT):
    """Set content of a table cell with mixed formatting. Clears existing."""
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = align
    p.paragraph_format.space_after  = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    for item in segments:
        if isinstance(item, str):
            _run(p, item)
        elif len(item) == 2:
            _run(p, item[0], bold=item[1])
        else:
            _run(p, item[0], bold=item[1], underline=item[2])


def _cell_plain(cell, text, bold=False, align=LEFT):
    _cell(cell, [(text, bold)])


# ---------------------------------------------------------------------------
# LETTER 1 — EXTENSION NOTICE
# ---------------------------------------------------------------------------

def generate_extension_notice(row: dict, letter_number: str, letter_date: date,
                               output_dir: str) -> str:
    doc = _new_doc()

    validity  = row['validity_date']
    ext_valid = extended_validity(validity)
    deadline  = submission_deadline(validity)
    amt_int   = row['bgb_amount_int']
    amt_words = amount_to_words(amt_int)

    # ------------------------------------------------------------------
    # 1. Contact block (left-aligned, "SPEED POST/ EMAIL" bold)
    # ------------------------------------------------------------------
    _para(doc, "SPEED POST/ EMAIL", bold=True)
    _para(doc, "Tele Mil : 6021")
    _para(doc, "Tele Civil : 0291-2515119")
    _para(doc, "Fax No 0291-2515113/2511519")
    _para(doc, "Email : cezjp2-mes@nic.in")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 2. HQ address
    # ------------------------------------------------------------------
    _para(doc, "Headquarters")
    _para(doc, "Chief Engineer Jodhpur Zone")
    _para(doc, "PIN-900066")
    _para(doc, "c/o 56 APO")

    # ------------------------------------------------------------------
    # 3. Letter number line
    # ------------------------------------------------------------------
    _para(doc, f"{letter_number}    {format_date(letter_date)}")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 4. Contractor address
    # ------------------------------------------------------------------
    _para(doc, row['contractor_name'])
    # Split contractor address on commas for multi-line look
    addr_parts = [p.strip() for p in row['contractor_address'].split(',') if p.strip()]
    for part in addr_parts:
        _para(doc, part)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 5. Subject — centered, bold, ALL CAPS
    # ------------------------------------------------------------------
    subject = (f"CA NO : {row['ca_number'].upper()} : "
               f"{row['work_name'].upper()}")
    _para(doc, subject, bold=True, align=CENTER)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 6. Salutation
    # ------------------------------------------------------------------
    _para(doc, "Dear Sir(s),", align=LEFT)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 7. Para 1 — bold: BGB number through amount words; also expiry date
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("1.\t", False),
        ("The Extended validity date in respect of Bank Guarantee Bond bearing No ", False),
        (f"{row['bgb_number']} dt {row['bgb_date_str']} of Rs {amt_int}/- "
         f"({amt_words} only)", True),
        (f" against Performance Security Deposit executed by "
         f"{row['bank_name']}, {row['bank_address']}"
         f" in your favour expires on ", False),
        (f"{format_date(validity)}.", True),
    ])

    # ------------------------------------------------------------------
    # 8. Para 2 — bold: extended date and deadline date
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("2.\t", False),
        ("It is requested that the validity of above Bank Guarantee Bonds may please be got "
         "extended from your bankers and sent along with six additional copies for the "
         "further period upto ", False),
        (f"{format_date(ext_valid)}", True),
        (" so as to reach to this HQ latest by ", False),
        (f"{format_date(deadline)}.", True),
    ])
    _para(doc, "")

    # ------------------------------------------------------------------
    # 9. Sign-off
    # ------------------------------------------------------------------
    _para(doc, "Yours faithfully,")
    _para(doc, "")
    _para(doc, "")
    _para(doc, "(MK Meena)")
    _para(doc, "EE (QS&C)")
    _para(doc, "Dy Dir (Contracts)")
    _para(doc, "for Chief Engineer")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 10. Copy to header (bold + underlined)
    # ------------------------------------------------------------------
    _para(doc, "Copy to:-", bold=True, underline=True)
    _para(doc, "SPEED POST", bold=True)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 11. Two-column copy-to (borderless table)
    #     Left: bank manager address | Right: numbered notes
    # ------------------------------------------------------------------
    tbl = _no_border_table(doc, 5, 2)
    # Column widths: ~3.3" left, ~4.0" right
    for r in tbl.rows:
        r.cells[0].width = Inches(3.3)
        r.cells[1].width = Inches(4.0)

    # Row 0
    _cell_plain(tbl.cell(0, 0), "The Manager")
    _cell(tbl.cell(0, 1), [
        ("1.\tIn case the contractor fails to extend the above bond by ", False),
        (f"{format_date(validity)}", False),
        (" then this letter may please be\n\ttreated as ", False),
        ("ENCASHMENT NOTICE", True, True),
    ])

    # Row 1
    _cell_plain(tbl.cell(1, 0), f"HDFC Bank Ltd")
    _cell_plain(tbl.cell(1, 1), "")

    # Row 2
    bank_addr_lines = [p.strip() for p in row['bank_address'].split(',') if p.strip()]
    _cell_plain(tbl.cell(2, 0), "\n".join(bank_addr_lines))
    _cell_plain(tbl.cell(2, 1),
                "2.\tPlease ensure that a copy of your letter\n"
                "\tendorse to regional office.")

    # Row 3
    _cell_plain(tbl.cell(3, 0), "")
    _cell_plain(tbl.cell(3, 1),
                "3.\tConfirmation be also given that signatory is\n"
                "\tauthorised to sign such BGB's and bind the Guarantor.")

    # Row 4 — blank spacer
    _cell_plain(tbl.cell(4, 0), "")
    _cell_plain(tbl.cell(4, 1), "")

    # ------------------------------------------------------------------
    # 12. Rest of copy-to list
    # ------------------------------------------------------------------
    _para(doc, "2.  PCDA SC Pune")
    _para(doc, "3.  CWE (A) Ahmedabad")
    _para(doc, f"4.  GE (A) {row['ge_location']}")
    _para(doc, f"5.  AO GE (A) {row['ge_location']}")
    _para(doc, "")
    _para(doc, "Internal :-", bold=True)
    _para(doc, "BGB Folder")

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------
    fname = build_output_filename(row['contractor_name'], row['ca_number'], "Extension")
    fpath = os.path.join(output_dir, fname)
    doc.save(fpath)
    return fpath


# ---------------------------------------------------------------------------
# LETTER 2 — ENCASHMENT NOTICE
# ---------------------------------------------------------------------------

def generate_encashment_notice(row: dict, letter_number: str, letter_date: date,
                                prev_letter_number: str, prev_letter_date: date,
                                output_dir: str) -> str:
    doc = _new_doc()

    validity  = row['validity_date']
    amt_int   = row['bgb_amount_int']
    amt_words = amount_to_words(amt_int)
    amt_indian = indian_amount_format(amt_int)   # e.g. "8,25,000.00"
    ge_loc    = row['ge_location']

    # ------------------------------------------------------------------
    # 1. Contact block ("E-Mail/SPEED POST" bold)
    # ------------------------------------------------------------------
    _para(doc, "E-Mail/SPEED POST", bold=True)
    _para(doc, "Tele Mil : 6021")
    _para(doc, "Tele Civil : 0291-2515119")
    _para(doc, "Fax No 0291-2515113/2511519")
    _para(doc, "Email : cezjp2-mes@nic.in")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 2. HQ address
    # ------------------------------------------------------------------
    _para(doc, "Headquarters")
    _para(doc, "Chief Engineer Jodhpur Zone")
    _para(doc, "PIN-900066")
    _para(doc, "c/o 56 APO")

    # ------------------------------------------------------------------
    # 3. Letter number line
    # ------------------------------------------------------------------
    _para(doc, f"{letter_number}    {format_date(letter_date)}")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 4. Addressee — Bank (each line separate)
    # ------------------------------------------------------------------
    _para(doc, "The Branch Manager")
    _para(doc, row['bank_name'])
    bank_addr_lines = [p.strip() for p in row['bank_address'].split(',') if p.strip()]
    for line in bank_addr_lines:
        _para(doc, line)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 5. Subject — centered, bold, ALL CAPS
    # ------------------------------------------------------------------
    subject = (f"CA NO : {row['ca_number'].upper()} : "
               f"{row['work_name'].upper()} : "
               f"{row['contractor_name'].upper()}")
    _para(doc, subject, bold=True, align=CENTER)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 6. Salutation
    # ------------------------------------------------------------------
    _para(doc, "Dear Sir (s),", align=LEFT)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 7. Para F — reference to previous letter
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("F-1.\t", False),
        ("Refer This HQ letter No. ", False),
        (f"{prev_letter_number} dt {format_date(prev_letter_date)}", False),
        (".", False),
    ])

    # ------------------------------------------------------------------
    # 8. Para 2 — bold: BGB details, CA number, contractor name, validity
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("2.\t", False),
        ("Whereas you have furnished a Bank Guarantee Bond bearing ", False),
        (f"{row['bgb_number']} dt {row['bgb_date_str']} for "
         f"Rs {amt_indian} (Rs. {amt_words} only) against Retention Money", True),
        (" executed against Contract Agreement No ", False),
        (f"{row['ca_number'].upper()}", True),
        (" made between President of India and M/s ", False),
        (f"{row['contractor_name'].upper()}", True),
        (" which is valid upto ", False),
        (f"{format_date(validity)}", True),
        (".", False),
    ])

    # ------------------------------------------------------------------
    # 9. Para 3 — fixed legal text
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("3.\t", False),
        ("And whereas it is expressly provided in the aforesaid guarantee bond that you "
         "undertake to pay the amount due and payable under the guarantee without any demur "
         "merely on a demand from the Govt, stating that the amount claimed is due by way of "
         "loss or damage caused to or suffered or would be caused to or suffered by the Govt. "
         "by reasons of any breach by the Contractor of any of the terms and conditions "
         "contained in the said contract agreement or by the reasons of contractor's failure "
         "to perform the said agreement.", False),
    ])

    # ------------------------------------------------------------------
    # 10. Para 4
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("4.\t", False),
        ("And whereas there has been a breach by the contractor of the terms and conditions "
         "of the agreement and amount claimed is due by way of loss or damage caused to or "
         "suffered or would be caused or suffered by the Govt.", False),
    ])

    # ------------------------------------------------------------------
    # 11. Para 5 — demand notice, bold amount
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("5.\t", False),
        ("I on behalf of the President of India, hereby serve you with this notice of demand "
         "against the aforesaid Bank Guarantee Bond for the sum of ", False),
        (f"Rs {amt_indian} (Rs. {amt_words} only)", True),
        (" on account of loss/damage caused to or suffered by the Govt.", False),
    ])

    # ------------------------------------------------------------------
    # 12. Para 6
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("6.\t", False),
        ("It is requested that the aforesaid sum as demanded be paid immediately.", False),
    ])

    # ------------------------------------------------------------------
    # 13. Para 7
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("7.\t", False),
        ("The amount may be remitted in the form of treasury challan in favour of ", False),
        (f"GE (A) {ge_loc}", True),
        (".", False),
    ])

    # ------------------------------------------------------------------
    # 14. Para 8a — bank details table
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("8.\t", False),
        ("The amount may be remitted in the form of treasury challan in favour of ", False),
        (f"GE (A) {ge_loc}", True),
        (" as per details mentioned below :-", False),
    ])

    tbl = _bordered_table(doc, 5, 2)
    tbl.columns[0].width = Inches(2.0)
    tbl.columns[1].width = Inches(5.3)
    details = [
        ("Account No.", row['account_number']),
        ("Name",        row['account_name']),
        ("IFSC Code",   row['ifsc_code']),
        ("Bank",        row['ge_bank_name']),
        ("Branch",      row['branch_name']),
    ]
    for i, (label, val) in enumerate(details):
        _cell_plain(tbl.cell(i, 0), label, bold=True)
        _cell_plain(tbl.cell(i, 1), val)

    _para(doc, "")

    # ------------------------------------------------------------------
    # 15. Para 8b — cancellation clause
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("8.\t", False),
        ("This NOTICE shall be TREATED AS cancelled if the above BGB has been extended "
         "within validity period.", False),
    ])

    # ------------------------------------------------------------------
    # 16. Para 9
    # ------------------------------------------------------------------
    _mixed(doc, [
        ("9.\t", False),
        ("Please ack receipt.", False),
    ])
    _para(doc, "")

    # ------------------------------------------------------------------
    # 17. Sign-off
    # ------------------------------------------------------------------
    _para(doc, "Yours faithfully,")
    _para(doc, "")
    _para(doc, "")
    _para(doc, "(Rajvinder Singh, VSM)")
    _para(doc, "Brig")
    _para(doc, "Chief Engineer")
    _para(doc, "Chief Engineer Jodhpur Zone")
    _para(doc, "FOR AND ON BEHALF OF THE PRESIDENT OF INDIA")
    _para(doc, "")

    # ------------------------------------------------------------------
    # 18. Copy to header (bold + underlined)
    # ------------------------------------------------------------------
    _para(doc, "Copy to:-", bold=True, underline=True)
    _para(doc, "SPEED POST/ e-Mail", bold=True)
    _para(doc, "")

    # ------------------------------------------------------------------
    # 19. Two-column copy-to (borderless table)
    #     Layout from PDF: left=name/address, right="- note"
    # ------------------------------------------------------------------
    tbl2 = _no_border_table(doc, 7, 3)
    # 3 cols: address | dash | note
    for r in tbl2.rows:
        r.cells[0].width = Inches(3.0)
        r.cells[1].width = Inches(0.2)
        r.cells[2].width = Inches(4.1)

    copy_rows = [
        # (left_lines, right_note)
        (
            ["The Regional Manager",
             f"HDFC Bank Ltd, GK Tower,",
             "2nd floor, Air port Road,",
             "Near Panch Batti Circle,",
             "Ratanada, Jodhpur"],
            "For info and necessary action please"
        ),
        (
            [row['contractor_name']] +
            [p.strip() for p in row['contractor_address'].split(',') if p.strip()],
            "Please extend the above BGB at the earliest."
        ),
        (["PCDA SC Pune"],                  ""),
        (["CWE (A) Ahmedabad"],             ""),
        ([f"GE (A) {ge_loc}"],
         "In case the above amount is not received by due date please inform "
         "this HQ through FAX/e-mail."),
        ([f"AO GE (A) {ge_loc}"],           ""),
        (["BGB Folder"],                    ""),
    ]

    for i, (left_lines, right_note) in enumerate(copy_rows):
        _cell_plain(tbl2.cell(i, 0), "\n".join(left_lines))
        _cell_plain(tbl2.cell(i, 1), "-" if right_note else "")
        _cell_plain(tbl2.cell(i, 2), right_note)

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------
    fname = build_output_filename(row['contractor_name'], row['ca_number'], "Encashment")
    fpath = os.path.join(output_dir, fname)
    doc.save(fpath)
    return fpath
